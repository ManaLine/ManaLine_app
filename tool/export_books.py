#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""One spreadsheet per business: people split by role, and one row per loan.

    python tool/export_books.py                 # to a temp folder, path printed
    python tool/export_books.py --out D:\\books  # somewhere you choose

Needs MANA_DB_URL -- the same libpq connection string tool/run_sql_tests.ps1
takes, and for the same reason it is not in this repo. See that script's
header: run.ps1.txt carries the anon key because it ships inside every APK
anyway, and a database password does not.

    $env:MANA_DB_URL = "postgresql://postgres:<password>@db.<ref>.supabase.co:5432/postgres"

WHY THIS EXISTS. Asked for on 2026-09-17: "give me entire list users in the
app till now - registered user with registered details from data base and
their loans - i want them for testing in an excell sheet, separate user
according to roles and one row per loan and one file for a business."

WHY IT IS READ-ONLY, AND STAYS THAT WAY. It runs two SELECTs and writes files.
It has no UPDATE, no INSERT and no DDL, and the session it opens is read-only
so that stays true even if somebody edits a query carelessly -- the same
runtime guard tool/run_sql_tests.ps1 uses, and for the reason recorded there:
static inspection of a query cannot prove what a function it calls will do.
"""
import argparse
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from urllib.parse import unquote, urlsplit

# --------------------------------------------------------------------------
# The two queries. Kept here rather than in supabase/tests/ because those are
# ASSERTIONS about the database and these are a report of it -- and because
# the runner over there refuses anything that is not marked scratch or
# production, which this is neither.
# --------------------------------------------------------------------------

PEOPLE_SQL = """
SELECT b.business_name, bm.role::text AS role, p.mlid, p.full_name,
       p.father_husband_name,
       CASE p.gender_digit WHEN '1' THEN 'Male' WHEN '2' THEN 'Female'
                           WHEN '3' THEN 'Others' ELSE p.gender_digit END AS gender,
       p.mobile_number,
       -- THE FULL AADHAAR IS NOT IN THIS DATABASE. persons.aadhaar_number is
       -- NULL for every row; only a hash and the last four digits are kept.
       -- Exporting the hash would be exporting a secret that identifies
       -- nobody, so the last four is what goes out.
       p.aadhaar_last4,
       p.mlid_type::text AS id_type, p.dob,
       p.preferred_language::text AS language,
       p.account_status::text AS account_status,
       p.profile_status::text AS profile_status,
       bm.membership_status::text AS membership_status,
       bm.onboarding_method::text AS onboarding_method,
       bm.verification_status::text AS verification_status,
       bm.joined_at, bm.created_at AS member_since,
       COALESCE(loc.village_town_name, '') AS village,
       COALESCE(pa.door_no, '')       AS door_no,
       COALESCE(pa.area_locality, '') AS area_locality,
       COALESCE(pa.pin_code, '')      AS pin_code,
       COALESCE(pa.mandal, '')        AS mandal,
       COALESCE(pa.district, '')      AS district,
       COALESCE(pa.state, '')         AS state
  FROM business_members bm
  JOIN businesses b ON b.business_id = bm.business_id
  JOIN persons p    ON p.person_id  = bm.person_id
  LEFT JOIN person_addresses pa ON pa.person_id = p.person_id AND pa.is_current
  LEFT JOIN locations loc ON loc.location_id = pa.village_id
 ORDER BY b.business_name, bm.role, p.full_name
"""

# LEFT JOINs on customers and persons, deliberately. An INNER JOIN here would
# DROP a loan whose customer row is missing -- and a Customer membership with
# no customers row is a real state this project has already found once, on
# 2026-09-17. A silently shorter export is worse than a row with blanks in it.
LOANS_SQL = """
SELECT b.business_name, l.loan_number,
       cp.mlid AS customer_mlid, cp.full_name AS customer_name,
       cp.father_husband_name AS customer_care_of,
       cp.mobile_number AS customer_mobile,
       COALESCE(loc.village_town_name,'') AS customer_village,
       l.repayment_amount, l.interest_amount, l.processing_fee, l.amount_given,
       l.repayment_type::text AS repayment_type, l.duration_value,
       l.installment_amount, l.grace_period_days, l.remaining_balance,
       l.effective_date, l.issue_business_date,
       l.loan_status::text AS loan_status, l.is_pre_existing,
       ap.full_name AS collection_agent, ap.mlid AS agent_mlid,
       (SELECT count(*) FROM collections c
         WHERE c.loan_id = l.loan_id AND c.deleted_at IS NULL) AS collections_count,
       (SELECT max(c.business_date) FROM collections c
         WHERE c.loan_id = l.loan_id AND c.deleted_at IS NULL) AS last_collection,
       (SELECT COALESCE(sum(c.collected_amount), 0) FROM collections c
         WHERE c.loan_id = l.loan_id AND c.deleted_at IS NULL) AS collected_total,
       l.closed_at, l.deleted_at, l.created_at
  FROM loans l
  JOIN businesses b ON b.business_id = l.business_id
  LEFT JOIN customers cu ON cu.customer_id = l.customer_id
  LEFT JOIN persons cp   ON cp.person_id = cu.person_id
  LEFT JOIN person_addresses pa ON pa.person_id = cp.person_id AND pa.is_current
  LEFT JOIN locations loc ON loc.location_id = pa.village_id
  LEFT JOIN business_members abm ON abm.membership_id = l.collection_agent_membership_id
  LEFT JOIN persons ap ON ap.person_id = abm.person_id
 ORDER BY b.business_name, cp.full_name, l.effective_date
"""

FONT = 'Arial'
ROLES = ['Owner', 'Agent', 'Investor', 'Customer']

DATE_KEYS = {'dob', 'joined_at', 'member_since', 'effective_date',
             'issue_business_date', 'last_collection', 'closed_at',
             'deleted_at', 'created_at'}
MONEY_KEYS = {'repayment_amount', 'interest_amount', 'processing_fee',
              'amount_given', 'installment_amount', 'remaining_balance',
              'collected_total'}
INT_KEYS = {'duration_value', 'grace_period_days', 'collections_count'}

# The file is named for the business, so repeating it on 64 rows is a column
# nobody reads.
DROP = {'business_name'}


def find_psql():
    found = shutil.which('psql')
    if found:
        return found
    for guess in (r'C:\Program Files\PostgreSQL\18\bin\psql.exe',
                  r'C:\Program Files\PostgreSQL\17\bin\psql.exe',
                  '/usr/bin/psql', '/usr/local/bin/psql'):
        if os.path.exists(guess):
            return guess
    sys.exit('psql was not found on PATH. Install the PostgreSQL client tools.')


def query(psql, conn, sql):
    """Run one SELECT and return its rows.

    THE PASSWORD GOES THROUGH PGPASSWORD, never the command line. Two reasons,
    both learned here and recorded in tool/run_sql_tests.ps1: a command line is
    readable by any other process on the machine for the life of the call, and
    a password containing a `$` once broke libpq's URI parsing badly enough
    that libpq printed the mis-parsed HOST -- with the password inside it -- to
    stderr. A credential that only leaks when something goes wrong leaks
    exactly when somebody is reading the output.
    """
    env = dict(os.environ)
    env['PGPASSWORD'] = conn['password']
    # Read-only for the session. A report has no business writing, and saying
    # so at runtime is stronger than saying so in a comment.
    env['PGOPTIONS'] = '-c default_transaction_read_only=on'
    wrapped = 'SELECT COALESCE(json_agg(t), \'[]\'::json) FROM (%s) t' % sql
    proc = subprocess.run(
        [psql, '--no-psqlrc', '-t', '-A', '-v', 'ON_ERROR_STOP=1',
         '-h', conn['host'], '-p', str(conn['port']),
         '-U', conn['user'], '-d', conn['dbname'], '-w', '-c', wrapped],
        env=env, capture_output=True, text=True, encoding='utf-8')
    if proc.returncode != 0:
        sys.exit('psql failed:\n' + (proc.stderr or '').strip())
    return json.loads(proc.stdout.strip() or '[]')


def parse_conn(url):
    parts = urlsplit(url)
    if parts.scheme not in ('postgresql', 'postgres'):
        sys.exit('MANA_DB_URL must be a postgresql:// URI.')
    user = unquote(parts.username or '')
    if not user or not parts.hostname:
        sys.exit('MANA_DB_URL is missing a host or a user.')
    return {
        'user': user,
        # Percent-decoded: a password correctly written as %24 must reach
        # libpq as `$`.
        'password': unquote(parts.password or ''),
        'host': parts.hostname,
        'port': parts.port or 5432,
        'dbname': (parts.path or '/postgres').lstrip('/') or 'postgres',
    }


def coerce(key, v):
    if v is None:
        return ''
    if key in DATE_KEYS and isinstance(v, str) and v:
        for fmt in ('%Y-%m-%dT%H:%M:%S.%f', '%Y-%m-%dT%H:%M:%S',
                    '%Y-%m-%d %H:%M:%S.%f', '%Y-%m-%d %H:%M:%S', '%Y-%m-%d'):
            try:
                d = datetime.datetime.strptime(v, fmt)
                return d.date() if fmt == '%Y-%m-%d' else d
            except ValueError:
                continue
        return v
    if key in MONEY_KEYS:
        try:
            return int(round(float(v)))
        except (TypeError, ValueError):
            return v
    if key in INT_KEYS:
        try:
            return int(v)
        except (TypeError, ValueError):
            return v
    return v


def header(s):
    special = {'mlid': 'MLID', 'dob': 'DOB', 'id': 'ID', 'pin': 'PIN'}
    return ' '.join(special.get(w, w.capitalize()) for w in s.split('_'))


def safe_sheet(name):
    return re.sub(r'[\\/*?:\[\]]', '-', name).strip()[:31]


def build(people, loans, outdir):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill('solid', fgColor='1F4E79')
    head_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
    body_font = Font(name=FONT, size=10)
    note_font = Font(name=FONT, size=10, italic=True, color='555555')
    bold_font = Font(name=FONT, bold=True, size=10)

    def write_sheet(ws, rows, keys):
        ws.append([header(k) for k in keys])
        for c in ws[1]:
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 28
        for r in rows:
            ws.append([coerce(k, r.get(k)) for k in keys])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
                k = keys[c.column - 1]
                if k in MONEY_KEYS:
                    c.number_format = '#,##0'
                elif k in DATE_KEYS:
                    c.number_format = 'yyyy-mm-dd'
        for i, k in enumerate(keys, start=1):
            widest = max([len(header(k))] +
                         [len(str(coerce(k, r.get(k)))) for r in rows] or [0])
            ws.column_dimensions[get_column_letter(i)].width = \
                min(max(widest + 2, 10), 38)
        ws.freeze_panes = 'A2'
        if rows:
            ws.auto_filter.ref = 'A1:%s%d' % (
                get_column_letter(len(keys)), len(rows) + 1)

    if not people:
        sys.exit('No rows came back. Is MANA_DB_URL pointing at the right database?')

    businesses = sorted({r['business_name'] for r in people} |
                        {r['business_name'] for r in loans})
    person_keys = [k for k in people[0].keys() if k not in DROP and k != 'role']
    loan_keys = [k for k in (loans[0].keys() if loans else []) if k not in DROP]
    made = []

    for b in businesses:
        wb = Workbook()
        summary = wb.active
        summary.title = 'Summary'
        counts = {}

        for role in ROLES:
            rows = sorted(
                (r for r in people
                 if r['business_name'] == b and r['role'] == role),
                key=lambda r: (r.get('full_name') or '').lower())
            name = safe_sheet(role + 's')
            write_sheet(wb.create_sheet(name), rows, person_keys)
            counts[name] = len(rows)

        lrows = sorted((r for r in loans if r['business_name'] == b),
                       key=lambda r: ((r.get('customer_name') or '').lower(),
                                      r.get('effective_date') or ''))
        write_sheet(wb.create_sheet('Loans'), lrows,
                    loan_keys or ['loan_number'])
        counts['Loans'] = len(lrows)

        summary['A1'] = b
        summary['A1'].font = Font(name=FONT, bold=True, size=13)
        summary['A3'] = 'Exported'
        summary['B3'] = datetime.datetime.now().strftime('%Y-%m-%d %H:%M')
        summary['A4'] = 'Source'
        summary['B4'] = 'MANA LINE database, tool/export_books.py'
        for cell in ('A6', 'B6', 'C6'):
            summary[cell].font = head_font
            summary[cell].fill = head_fill
        summary['A6'], summary['B6'], summary['C6'] = \
            'Sheet', 'Rows Now', 'Rows At Export'

        row = 7
        for name in ['Owners', 'Agents', 'Investors', 'Customers', 'Loans']:
            summary.cell(row=row, column=1, value=name).font = body_font
            # A live formula AND a frozen count, which is not redundancy.
            #
            # openpyxl writes a formula with no cached value, so "Rows Now"
            # reads blank until the file is opened in Excel or Sheets. The
            # frozen count is readable anywhere -- and afterwards it earns its
            # place, because this is a workbook somebody deletes rows in while
            # testing: one column follows what is left, the other says what
            # came out of the database, and a difference between them is
            # information rather than a fault.
            summary.cell(row=row, column=2,
                         value="=COUNTA('%s'!A2:A100000)" % name).font = body_font
            summary.cell(row=row, column=3, value=counts[name]).font = body_font
            row += 1
        summary.cell(row=row, column=1, value='Total people').font = bold_font
        summary.cell(row=row, column=2, value='=SUM(B7:B10)').font = bold_font
        summary.cell(row=row, column=3,
                     value=sum(counts[n] for n in
                               ['Owners', 'Agents', 'Investors',
                                'Customers'])).font = bold_font

        notes = [
            '',
            'NOTES',
            'One row per person per ROLE. Somebody who is both an Agent and a '
            'Customer of this business appears on both sheets - that is two '
            'memberships, which is how the database holds it.',
            'Loans: one row per loan, including closed, deleted and '
            'pre-existing ones. Collections Count and Collected Total exclude '
            'deleted collections.',
            'Aadhaar: the FULL number is not stored anywhere in this database '
            '- only a hash and the last four digits. Aadhaar Last4 is '
            'everything there is.',
            'Money is whole rupees. Every money column in the database is '
            'numeric(_,0); paise cannot be stored.',
            '"Rows Now" is a live formula and reads blank until this file is '
            'opened in a spreadsheet. "Rows At Export" is frozen and always '
            'readable.',
        ]
        row += 2
        for line in notes:
            c = summary.cell(row=row, column=1, value=line)
            c.font = bold_font if line == 'NOTES' else note_font
            c.alignment = Alignment(wrap_text=True, vertical='top')
            summary.merge_cells(start_row=row, start_column=1,
                                end_row=row, end_column=6)
            if line and line != 'NOTES':
                summary.row_dimensions[row].height = 30
            row += 1

        summary.column_dimensions['A'].width = 26
        summary.column_dimensions['B'].width = 14
        summary.column_dimensions['C'].width = 16
        for col in 'DEF':
            summary.column_dimensions[col].width = 18

        path = os.path.join(
            outdir, re.sub(r'[^\w.-]+', '_', b).strip('_') + '.xlsx')
        wb.save(path)
        made.append((path, counts))
    return made


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--out', help='Where to write the workbooks. Defaults to a '
                                  'new folder under the system temp directory.')
    args = ap.parse_args()

    url = os.environ.get('MANA_DB_URL')
    if not url:
        sys.exit(
            'MANA_DB_URL is not set, so there is nothing to export from.\n'
            '  $env:MANA_DB_URL = "postgresql://postgres:<password>'
            '@db.<ref>.supabase.co:5432/postgres"\n'
            '  python tool/export_books.py')

    outdir = args.out or os.path.join(
        tempfile.gettempdir(),
        'mana_books_' + datetime.date.today().isoformat())

    # NOT INTO THE REPO, unless somebody insists by pointing --out at it and
    # meaning it. These files are a dump of real people's names, mobile
    # numbers and addresses; a working tree is one `git add -A` away from
    # publishing them, and that mistake is not recoverable from a public
    # remote. Defaulting elsewhere is the guard; this catches the slip.
    repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    if os.path.abspath(outdir).startswith(repo + os.sep):
        sys.exit(
            'Refusing to write inside the repository: %s\n'
            'These workbooks hold names, mobile numbers and addresses, and a '
            'working tree is one "git add -A" from publishing them.\n'
            'Pass --out somewhere outside the repo, or omit it for a temp '
            'folder.' % outdir)

    os.makedirs(outdir, exist_ok=True)
    psql = find_psql()
    conn = parse_conn(url)

    people = query(psql, conn, PEOPLE_SQL)
    loans = query(psql, conn, LOANS_SQL)

    made = build(people, loans, outdir)

    print('%d people, %d loans' % (len(people), len(loans)))
    for path, counts in made:
        print('  %-36s %s' % (os.path.basename(path), counts))
    print('\nWritten to %s' % outdir)


if __name__ == '__main__':
    main()
